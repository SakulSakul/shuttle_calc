"""build_styled_xlsx 단위 테스트 — 더미 데이터로 시트/수식 생성 확인."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from settlement import REQUIRED_COLS, build_settlement
from excel_format import build_styled_xlsx


def _dummy():
    rows = [
        ("스위스관광", "T1", "김", "A", "111", "중소"),
        ("신백승여행사", "T2", "이", "A", "111", "중소"),
        ("한진관광", "T3", "박", "A", "111", "중소"),
        ("스위스관광", "T4", "최", "B", "222", "중견"),
        ("한진관광", "T5", "정", "B", "222", "중견"),
    ]
    df = pd.DataFrame(rows, columns=REQUIRED_COLS)
    return build_settlement(df, 41935)


def main():
    up, final_df = _dummy()
    data = build_styled_xlsx(up, final_df, 41935)
    assert isinstance(data, (bytes, bytearray)) and len(data) > 0, "바이트 출력 아님"

    wb = load_workbook(BytesIO(data))
    sheets = wb.sheetnames
    assert sheets == ["실제_탑승인원_명단", "협력사별_지원금액_총계"], f"시트 구성 불일치: {sheets}"

    ws = wb["협력사별_지원금액_총계"]
    # 가정 셀 C2 = support_amount
    assert ws["C2"].value == 41935, f"C2 가정셀 값 불일치: {ws['C2'].value}"

    # 총계 행에 =SUM 수식이 있어야 함 (값이 아니라 수식)
    sum_formulas = [
        c.value for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=SUM(")
    ]
    assert sum_formulas, "총계 SUM 수식이 없음"

    # 데이터 금액 셀은 인원*$C$2 수식
    mult_formulas = [
        c.value for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and "*$C$2" in c.value
    ]
    assert mult_formulas, "금액 셀의 *$C$2 수식이 없음"

    # 운영사 하드코딩 없이 3개 운영사 헤더가 동적 생성됐는지
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    for op in ("스위스관광", "신백승여행사", "한진관광"):
        assert op in headers, f"운영사 헤더 누락: {op}"
        assert f"{op} 지원금액" in headers, f"운영사 지원금액 헤더 누락: {op}"

    # 명단 시트 메타 '총 N명'
    ws1 = wb["실제_탑승인원_명단"]
    assert ws1["A2"].value == "총 5명", f"명단 메타 불일치: {ws1['A2'].value}"

    print("PASS - 시트 2개 생성")
    print("PASS - C2 가정셀 =", ws["C2"].value)
    print(f"PASS - 총계 SUM 수식 {len(sum_formulas)}개")
    print(f"PASS - 금액 *$C$2 수식 {len(mult_formulas)}개")
    print("PASS - 3개 운영사 헤더 동적 생성")
    print("PASS - 명단 '총 5명' 메타")
    print("\nALL PASS")


if __name__ == "__main__":
    main()
