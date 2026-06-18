"""엑셀 출력 서식(표현) 레이어.

계산 로직(settlement.py)과 완전히 분리되어 있으며, 운영사 컬럼명을 하드코딩하지
않고 final_df 의 컬럼명을 보고 금액/인원/텍스트를 동적으로 판별한다.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ── 브랜드 색상 (신세계 레드 기준) ────────────────────────────────
RED = "FFC8102E"        # 헤더 배경 / 강조 글자
DARK = "FF1A1A1A"       # 타이틀 / 총계 상단선
WHITE = "FFFFFFFF"
ZEBRA = "FFF7F7F7"      # 짝수행 zebra 배경
BORDER_GREY = "FFD9D9D9"
TOTAL_BG = "FFFBEAED"   # 총계 행 배경
ASSUMP_BG = "FFFFF8E1"  # 1인당 지원금액 가정 셀 배경

FONT_NAME = "맑은 고딕"

# ── 숫자 서식 ────────────────────────────────────────────────────
FMT_CURRENCY = '₩#,##0'
FMT_PEOPLE = '#,##0"명"'

# ── 공통 스타일 객체 ─────────────────────────────────────────────
_thin = Side(style="thin", color=BORDER_GREY)
THIN_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
TOTAL_BORDER = Border(
    left=_thin, right=_thin, bottom=_thin,
    top=Side(style="medium", color=DARK),
)

HEADER_FILL = PatternFill("solid", fgColor=RED)
ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
TOTAL_FILL = PatternFill("solid", fgColor=TOTAL_BG)
ASSUMP_FILL = PatternFill("solid", fgColor=ASSUMP_BG)

INDEX_TEXT_COLS = ("협력회사명", "사업자등록번호")  # 좌측 정렬 텍스트
AMOUNT_SUFFIX = " 지원금액"


def _f(_name=FONT_NAME, **kw):
    """맑은 고딕 폰트 헬퍼 (폰트명을 한 곳에서 통일)."""
    return Font(name=FONT_NAME, **kw)


def _col_kind(colname):
    """컬럼명을 보고 (number_format, alignment) 를 동적으로 결정.

    - '지원금액' 으로 끝나면 통화(우측)
    - '기업규모' 는 가운데 텍스트
    - 협력회사명/사업자등록번호 는 좌측 텍스트
    - 그 외(운영사 인원 컬럼, '총 인원')는 인원("명", 우측)
    """
    if colname.endswith(AMOUNT_SUFFIX):
        return FMT_CURRENCY, "right", "amount"
    if colname == "기업규모":
        return None, "center", "text"
    if colname in INDEX_TEXT_COLS:
        return None, "left", "text"
    return FMT_PEOPLE, "right", "count"


def _amount_to_count_letter(cols):
    """금액 컬럼명 -> 대응하는 인원 컬럼의 엑셀 열 문자 매핑."""
    name_to_letter = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
    mapping = {}
    for c in cols:
        if not c.endswith(AMOUNT_SUFFIX):
            continue
        base = "총 인원" if c == "총 지원금액" else c[: -len(AMOUNT_SUFFIX)]
        if base in name_to_letter:
            mapping[c] = name_to_letter[base]
    return mapping


def _col_width(colname):
    if colname == "협력회사명":
        return 24
    if colname == "사업자등록번호":
        return 18
    if colname.endswith(AMOUNT_SUFFIX):
        return 16
    if colname in ("운영사", "탑승자", "태그ID"):
        return 14
    return 13


def _write_header(ws, header_row, columns):
    for ci, name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=ci, value=name)
        cell.fill = HEADER_FILL
        cell.font = _f(FONT_NAME, color=WHITE, bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(ci)].width = _col_width(name)
    ws.row_dimensions[header_row].height = 40


def _finish_sheet(ws):
    """격자선 숨김 + landscape + fitToWidth."""
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def _build_passenger_sheet(ws, unique_passengers):
    columns = list(unique_passengers.columns)
    ncols = len(columns)
    last_letter = get_column_letter(ncols)

    # 행1 타이틀
    ws.merge_cells(f"A1:{last_letter}1")
    t = ws.cell(row=1, column=1, value="실제 탑승인원 명단")
    t.font = _f(FONT_NAME, bold=True, size=16, color=DARK)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    # 행2 메타: 총 N명
    ws.merge_cells("A2:B2")
    m = ws.cell(row=2, column=1, value=f"총 {len(unique_passengers)}명")
    m.font = _f(FONT_NAME, bold=True, size=11, color=RED)
    m.alignment = Alignment(horizontal="left", vertical="center")

    header_row = 4
    _write_header(ws, header_row, columns)

    for ri in range(len(unique_passengers)):
        er = header_row + 1 + ri
        for ci, name in enumerate(columns, start=1):
            fmt, align, _ = _col_kind(name)
            # 명단 시트는 모두 텍스트(운영사/태그ID/탑승자 등) — 숫자서식 미적용
            cell = ws.cell(row=er, column=ci, value=unique_passengers.iloc[ri][name])
            cell.font = _f(FONT_NAME, size=10)
            halign = "center" if name == "기업규모" else "left"
            cell.alignment = Alignment(horizontal=halign, vertical="center")
            cell.border = THIN_BORDER
            if ri % 2 == 1:
                cell.fill = ZEBRA_FILL
        ws.row_dimensions[er].height = 22

    ws.freeze_panes = "A5"
    _finish_sheet(ws)


def _build_settlement_sheet(ws, final_df, support_amount):
    columns = list(final_df.columns)
    ncols = len(columns)
    last_letter = get_column_letter(ncols)
    amount_count_letter = _amount_to_count_letter(columns)

    # 행1 타이틀
    ws.merge_cells(f"A1:{last_letter}1")
    t = ws.cell(row=1, column=1, value="협력사별 지원금액 총계")
    t.font = _f(FONT_NAME, bold=True, size=16, color=DARK)
    t.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    # 행2 메타: A2:B2 라벨 + C2 가정 셀(1인당 지원금액)
    ws.merge_cells("A2:B2")
    lab = ws.cell(row=2, column=1, value="1인당 지원금액(원)")
    lab.font = _f(FONT_NAME, bold=True, size=11, color=DARK)
    lab.alignment = Alignment(horizontal="right", vertical="center")
    amt = ws.cell(row=2, column=3, value=support_amount)
    amt.fill = ASSUMP_FILL
    amt.font = _f(FONT_NAME, bold=True, size=11, color=DARK)
    amt.number_format = FMT_CURRENCY
    amt.alignment = Alignment(horizontal="right", vertical="center")
    amt.border = THIN_BORDER

    header_row = 4
    _write_header(ws, header_row, columns)

    # 데이터 행은 총계 행을 제외 (마지막 행이 '총계')
    data = final_df.iloc[:-1]
    num_data = len(data)
    first_data_row = header_row + 1            # 5
    last_data_row = header_row + num_data       # 5 + num_data - 1
    total_row = first_data_row + num_data       # 총계 행

    for ri in range(num_data):
        er = first_data_row + ri
        for ci, name in enumerate(columns, start=1):
            fmt, align, kind = _col_kind(name)
            letter = get_column_letter(ci)
            if kind == "amount" and name in amount_count_letter:
                # 금액 = 인원 * 가정셀(C2)  → 가정 셀 변경 시 자동 재계산
                clet = amount_count_letter[name]
                cell = ws.cell(row=er, column=ci, value=f"={clet}{er}*$C$2")
            else:
                cell = ws.cell(row=er, column=ci, value=data.iloc[ri][name])
            cell.font = _f(FONT_NAME, size=10)
            if fmt:
                cell.number_format = fmt
            cell.alignment = Alignment(horizontal=align, vertical="center")
            cell.border = THIN_BORDER
            if ri % 2 == 1:
                cell.fill = ZEBRA_FILL
        ws.row_dimensions[er].height = 22

    # 총계 행 — 숫자 컬럼은 =SUM(범위) 수식
    for ci, name in enumerate(columns, start=1):
        fmt, align, kind = _col_kind(name)
        letter = get_column_letter(ci)
        if kind in ("amount", "count"):
            rng = f"{letter}{first_data_row}:{letter}{last_data_row}"
            cell = ws.cell(row=total_row, column=ci, value=f"=SUM({rng})")
        else:
            # 텍스트 컬럼: final_df 총계 행 값('총계' / '-')
            cell = ws.cell(row=total_row, column=ci, value=final_df.iloc[-1][name])
        is_amount = kind == "amount"
        cell.font = _f(FONT_NAME, bold=True, size=10, color=(RED if is_amount else DARK))
        if fmt:
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal=align, vertical="center")
        cell.fill = TOTAL_FILL
        cell.border = TOTAL_BORDER
    ws.row_dimensions[total_row].height = 24

    ws.freeze_panes = "A5"
    _finish_sheet(ws)


def build_styled_xlsx(unique_passengers, final_df, support_amount):
    """2개 시트(명단/정산)를 서식과 함께 작성해 xlsx 바이트로 반환."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "실제_탑승인원_명단"
    _build_passenger_sheet(ws1, unique_passengers)

    ws2 = wb.create_sheet("협력사별_지원금액_총계")
    _build_settlement_sheet(ws2, final_df, support_amount)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
